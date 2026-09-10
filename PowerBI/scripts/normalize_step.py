#!/usr/bin/env python3
"""Normalize the two STEP raw exports into lean, PII-free CSVs for Power BI.

Inputs (repo root):
    RegisterInterest.xlsx   - one row per interest registration
    CourseApplicant.xlsx    - one row per course application (sign-up)

Outputs (repo root):
    STEP_RegisterInterest.csv
    STEP_CourseApplicant.csv

Design notes:
- ALL personal data (names, emails, phone numbers, employer/company details,
  applicant background) is dropped. Only aggregate/analytical fields are kept.
- Dates are written ISO (yyyy-mm-dd) so Power Query parses them reliably.
- 'RegDate' is the event date used to relate each table to the Calendar:
    * RegisterInterest -> Date of Register of Interest
    * CourseApplicant  -> Registration time (fallback: Course start date)
"""
import csv
import datetime
import os

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))


def d(v):
    """Return an ISO yyyy-mm-dd string for a date/datetime cell, else ''."""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    # occasionally dates arrive as text
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def clean(v):
    return "" if v is None else str(v).strip()


def header_index(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}


# ---------------------------------------------------------------- Interest
def build_interest():
    wb = openpyxl.load_workbook(os.path.join(ROOT, "RegisterInterest.xlsx"),
                                read_only=True, data_only=True)
    ws = wb["Sheet1"]
    hi = header_index(ws)
    # the raw date header has a trailing space
    date_key = next(k for k in hi if k and k.strip() == "Date of Register of Interest")
    out = os.path.join(ROOT, "STEP_RegisterInterest.csv")
    n = 0
    with open(out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["RegDate", "Interested Course Code", "Interested Course Name",
                    "Course Category", "Course Type", "Area of Interest", "Sub-interest"])
        for row in ws.iter_rows(min_row=2, values_only=True):
            def g(name):
                return row[hi[name] - 1]
            reg = d(g(date_key))
            if reg == "" and clean(g("Interested Course Code")) == "":
                continue
            w.writerow([reg, clean(g("Interested Course Code")), clean(g("Interested Course Name")),
                        clean(g("Course Category")), clean(g("Course Type")),
                        clean(g("Area of interest")), clean(g("Sub-interest"))])
            n += 1
    wb.close()
    print(f"STEP_RegisterInterest.csv: {n} rows")
    return n


# --------------------------------------------------------------- Applicant
def build_applicant():
    wb = openpyxl.load_workbook(os.path.join(ROOT, "CourseApplicant.xlsx"),
                                read_only=True, data_only=True)
    ws = wb["Sheet1"]
    hi = header_index(ws)
    out = os.path.join(ROOT, "STEP_CourseApplicant.csv")
    n = 0
    with open(out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Application ID", "Course Category", "Course Type", "Course Code",
                    "Course Name", "Course Intake No", "Course Start Date", "Course End Date",
                    "Course Intake Status", "School", "Branch", "Application Status", "RegDate"])
        for row in ws.iter_rows(min_row=2, values_only=True):
            def g(name):
                return row[hi[name] - 1]
            appid = clean(g("Application ID"))
            if appid == "":
                continue
            reg = d(g("Registration time")) or d(g("Course start date"))
            w.writerow([appid, clean(g("Course category")), clean(g("Course type")),
                        clean(g("Course Code")), clean(g("Course Name")),
                        clean(g("Course intake No.")), d(g("Course start date")),
                        d(g("Course end date")), clean(g("Course intake status")),
                        clean(g("School")), clean(g("Branch")),
                        clean(g("Application status")), reg])
            n += 1
    wb.close()
    print(f"STEP_CourseApplicant.csv: {n} rows")
    return n


if __name__ == "__main__":
    build_interest()
    build_applicant()
